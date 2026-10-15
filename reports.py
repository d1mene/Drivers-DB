"""Отчёты: вывод на экран + сохранение в .txt или .xlsx.

Расчёты выполняются хранимыми процедурами:
get_fastest_trip(route, OUT min_time, OUT best_auto) - рекорд по маршруту
get_route_json_info(route) - автомобили в рейсе
"""
import json
import tkinter as tk
from datetime import datetime
from psycopg2 import errors
from tkinter import ttk, filedialog, messagebox
from db import error_text, unique_labels

DRIVER_SQL = ("SELECT a.num, COALESCE(p.last_name || ' ' || p.first_name, '—') "
              "FROM auto a LEFT JOIN auto_personal p ON p.id = a.personal_id WHERE a.id = %s")


def _routes(db, route_id):
    _, rows = db.query("SELECT id, name FROM routes ORDER BY name, id")
    pairs = unique_labels(rows)
    return [p for p in pairs if route_id is None or p[0] == route_id]


def report_record(db, route_id):
    headers = ["Маршрут", "Рекордное время, мин", "Автомобиль", "Водитель"]
    rows = []
    for rid, name in _routes(db, route_id):
        try:
            _, r = db.query("CALL get_fastest_trip(%s, NULL, NULL)", (rid,))
            minutes, auto_id = r[0]
            _, a = db.query(DRIVER_SQL, (auto_id,))
            rows.append((name, minutes, a[0][0], a[0][1]))
        except errors.RaiseException:
            rows.append((name, "нет завершённых рейсов", "—", "—"))
    return headers, rows, []


def report_active(db, route_id):
    headers = ["Маршрут", "Автомобилей в рейсе", "Номера автомобилей"]
    rows, total = [], 0
    for rid, name in _routes(db, route_id):
        _, r = db.query("SELECT get_route_json_info(%s)", (rid,))
        info = r[0][0]
        if isinstance(info, str):
            info = json.loads(info)
        cnt = info["active_cars_count"]
        total += cnt
        rows.append((name, cnt, ", ".join(info["cars"]) or "—"))
    return headers, rows, [f"Всего автомобилей в рейсе: {total}"]


REPORTS = {
    "record": dict(
        title="Рекордное время проезда по маршруту",
        text="Самое короткое время проезда по маршруту и автомобиль-рекордсмен.",
        func=report_record),
    "active": dict(
        title="Автомобили, находящиеся в рейсе",
        text="Количество и номера автомобилей, которые сейчас на маршруте "
             "(время прибытия в журнале не заполнено).",
        func=report_active),
}


def render_txt(title, info, headers, rows, footer):
    srows = [["" if v is None else str(v) for v in r] for r in rows]
    widths = [max([len(h)] + [len(r[i]) for r in srows]) for i, h in enumerate(headers)]
    line = lambda cells: " | ".join(c.ljust(w) for c, w in zip(cells, widths))
    out = [title, "=" * len(title), *info, "", line(headers),
           "-+-".join("-" * w for w in widths), *[line(r) for r in srows]]
    if not srows:
        out.append("(нет данных)")
    if footer:
        out += ["", *footer]
    return "\n".join(out) + "\n"


def save_xlsx(path, title, info, headers, rows, footer):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    r = 1
    ws.cell(r, 1, title).font = Font(bold=True, size=14)
    
    for text in info:
        r += 1
        ws.cell(r, 1, text)
    r += 2
    
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center")
        
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    for text in footer:
        r += 2
        ws.cell(r, 1, text).font = Font(bold=True)
    for c, h in enumerate(headers, 1):
        w = max([len(str(h))] + [len(str(row[c - 1])) for row in rows]) + 3
        ws.column_dimensions[get_column_letter(c)].width = min(w, 60)
    wb.save(path)


class ReportWindow(tk.Toplevel):
    ALL = "Все маршруты"

    def __init__(self, master, db, key):
        super().__init__(master)
        self.db, self.spec = db, REPORTS[key]
        self.key, self.result = key, None
        self.title("Отчёт: " + self.spec["title"])
        self.geometry("820x520")

        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")
        ttk.Label(top, text=self.spec["text"], wraplength=780).grid(
            row=0, column=0, columnspan=4, sticky="w", pady=(0, 8))
        ttk.Label(top, text="Маршрут:").grid(row=1, column=0, sticky="w")
        try:
            self.pairs = unique_labels(db.query("SELECT id, name FROM routes ORDER BY name, id")[1])
        except Exception as e:
            messagebox.showerror("Ошибка", error_text(e), parent=self)
            self.pairs = []
        self.route_var = tk.StringVar(value=self.ALL)
        ttk.Combobox(top, textvariable=self.route_var, state="readonly", width=48,
                     values=[self.ALL] + [l for _, l in self.pairs]
                     ).grid(row=1, column=1, columnspan=3, sticky="w", padx=6)

        ttk.Label(top, text="Формат файла:").grid(row=2, column=0, sticky="w", pady=8)
        self.fmt = tk.StringVar(value="txt")
        ttk.Radiobutton(top, text="Текст (.txt)", variable=self.fmt, value="txt").grid(row=2, column=1, sticky="w", padx=6)
        ttk.Radiobutton(top, text="Excel (.xlsx)", variable=self.fmt, value="xlsx").grid(row=2, column=2, sticky="w")

        btns = ttk.Frame(top)
        btns.grid(row=3, column=0, columnspan=4, sticky="w")
        ttk.Button(btns, text="Сформировать", command=self.generate).pack(side="left")
        ttk.Button(btns, text="Сохранить в файл…", command=self.save).pack(side="left", padx=6)
        ttk.Button(btns, text="Закрыть", command=self.destroy).pack(side="left")

        frame = ttk.Frame(self, padding=(10, 0, 10, 10))
        frame.pack(fill="both", expand=True)
        self.text = tk.Text(frame, wrap="none", font=("Courier New", 10), state="disabled")
        vs = ttk.Scrollbar(frame, command=self.text.yview)
        hs = ttk.Scrollbar(frame, orient="horizontal", command=self.text.xview)
        self.text.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.text.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

    def _route_id(self):
        label = self.route_var.get()
        for rid, lbl in self.pairs:
            if lbl == label:
                return rid
        return None

    def generate(self):
        try:
            headers, rows, footer = self.spec["func"](self.db, self._route_id())
        except Exception as e:
            messagebox.showerror("Ошибка формирования отчёта", error_text(e), parent=self)
            return
        info = [f"Сформирован: {datetime.now():%d.%m.%Y %H:%M:%S}",
                f"Маршрут: {self.route_var.get()}"]
        self.result = (self.spec["title"], info, headers, rows, footer)
        self.text.config(state="normal")
        self.text.delete("1.0", "end")
        self.text.insert("1.0", render_txt(*self.result))
        self.text.config(state="disabled")

    def save(self):
        if self.result is None:
            self.generate()
            if self.result is None:
                return
        ext = self.fmt.get()
        types = [("Текстовый файл", "*.txt")] if ext == "txt" else [("Excel", "*.xlsx")]
        path = filedialog.asksaveasfilename(
            parent=self, defaultextension="." + ext, filetypes=types,
            initialfile=f"report_{self.key}_{datetime.now():%Y%m%d_%H%M}.{ext}")
        if not path:
            return
        try:
            if ext == "txt":
                with open(path, "w", encoding="utf-8-sig") as f:
                    f.write(render_txt(*self.result))
            else:
                save_xlsx(path, *self.result)
        except ImportError:
            messagebox.showerror("Нет библиотеки",
                                 "Для .xlsx установите openpyxl:\npip install openpyxl", parent=self)
            return
        except Exception as e:
            messagebox.showerror("Не удалось сохранить файл", str(e), parent=self)
            return
        messagebox.showinfo("Готово", f"Отчёт сохранён:\n{path}", parent=self)
